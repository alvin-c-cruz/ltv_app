"""Verification for knocked-out contract presentation in the LTV Stocks workbook.

A contract knocks out the first day its closing price crosses the K/O barrier
(>= ko for an accumulator, <= ko for a decumulator). The hand-maintained
workbook marks that week's KO'd rows in a fixed way, and drops the contract
entirely from the following week's report once it is flagged `status='KO'` in
`tbl_stock_contract`. This pins all of it:

  1. `find_ko_day` -- isolated classification, no DB: which O:X column (if any)
     is the knock-out day, under the same skip rules as compute_status_flags
     (outside start/end, holidays, missing/zero prices).
  2. The rendered row -- synthetic records written into a real worksheet:
     KO-day price bold, the next O:X column holding a bold "KO" literal, every
     column after it blank and grey-filled, the whole row A:AJ in red, and the
     block's head count reduced by the number of KO'd rows.
  3. `compute_status_flags` stops circling a row at its knock-out day.
  4. `contract_records` excludes `status='KO'` contracts (live, against the
     real DB -- a structural assertion, no client figures embedded here).

Run: server/.venv/Scripts/python.exe scripts/verify_ko_presentation.py
"""
import os
import sys
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
SERVER = os.path.dirname(HERE)
sys.path.insert(0, SERVER)

import openpyxl

from ltv_app import create_app
from ltv_app.blueprints.database.views import get_db
from ltv_app.blueprints.ltv_stocks.legacy_port.excel_writer import (
    _GREY_FILL, _OX_COLS, _STATUS_COLS, _ko_offsets, _write_contracts,
    _write_status_grid,
)
from ltv_app.blueprints.ltv_stocks.legacy_port.report_data import (
    compute_status_flags, find_ko_day,
)

_RED = 'FFFF0000'
_ALL_ROW_COLS = ('A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L',
                 'M', 'N') + _OX_COLS + ('Z',) + _STATUS_COLS

ok = True


def check(label, actual, expected):
    global ok
    if actual == expected:
        print(f"  {label}: PASS")
    else:
        ok = False
        print(f"  {label}: FAIL  expected {expected!r}, got {actual!r}")


class _NeverHoliday:
    def is_holiday(self, d):
        return False


class _Holidays:
    def __init__(self, days):
        self.days = set(days)

    def is_holiday(self, d):
        return d in self.days


# A 10-day window shaped like week_dates(): prev Mon-Fri + current Mon-Fri.
DATES = [date(2026, 6, 1), date(2026, 6, 2), date(2026, 6, 3), date(2026, 6, 4),
         date(2026, 6, 5), date(2026, 6, 8), date(2026, 6, 9), date(2026, 6, 10),
         date(2026, 6, 11), date(2026, 6, 12)]
REPORT_DATE = DATES[-1]


def _rec(**kw):
    """A contract record with every field _write_contracts reads."""
    base = dict(code_ref=1, stock_name='Test Stock NO GTD', stock_name_plain='Test Stock',
                code='0001', bank_doc='TEST-1', shares='100 / 200', spot=100.0,
                strike=80.0, ko=120.0, start_date=DATES[0], end_date=date(2027, 6, 1),
                received=1, total=12, yahoo_ticker='0001.HK', next_date=date(2026, 7, 1),
                frequency='monthly', ccy_id='HKD')
    base.update(kw)
    return base


def _prices(mapping):
    return lambda code_ref, d: mapping.get(d)


def isolated_find_ko_day():
    print("Isolated find_ko_day checks:")
    wd = _NeverHoliday()

    # Accumulator shape (spot > strike): knocks out when closing >= ko.
    accu = _rec(spot=100.0, strike=80.0, ko=120.0)
    px = _prices({DATES[0]: 100.0, DATES[1]: 110.0, DATES[2]: 125.0,
                  DATES[3]: 130.0, DATES[4]: 90.0})
    check("accumulator KO on first breach", find_ko_day(accu, DATES, wd, px), 2)

    # Never breaches.
    px = _prices({d: 100.0 for d in DATES})
    check("no breach -> None", find_ko_day(accu, DATES, wd, px), None)

    # Decumulator shape (spot < strike): knocks out when closing <= ko.
    decu = _rec(spot=100.0, strike=130.0, ko=90.0)
    px = _prices({DATES[0]: 100.0, DATES[1]: 95.0, DATES[2]: 95.0,
                  DATES[3]: 89.0, DATES[4]: 80.0})
    check("decumulator KO on first breach", find_ko_day(decu, DATES, wd, px), 3)

    # An equal close counts as a breach on both sides (<= / >=).
    px = _prices({DATES[0]: 90.0})
    check("decumulator KO at exactly the barrier", find_ko_day(decu, DATES, wd, px), 0)
    px = _prices({DATES[0]: 120.0})
    check("accumulator KO at exactly the barrier", find_ko_day(accu, DATES, wd, px), 0)

    # Missing / zero / non-numeric prices are not breaches, they are gaps.
    px = _prices({DATES[0]: None, DATES[1]: 0, DATES[2]: 'Done', DATES[3]: 125.0})
    check("missing, zero and text prices skipped", find_ko_day(accu, DATES, wd, px), 3)

    # A day outside the contract's own start/end window is never a KO day,
    # even with a breaching price on it.
    windowed = _rec(spot=100.0, strike=80.0, ko=120.0,
                    start_date=DATES[3], end_date=DATES[4])
    px = _prices({d: 125.0 for d in DATES})
    check("breach before start_date ignored", find_ko_day(windowed, DATES, wd, px), 3)

    # Holidays are skipped even if the lookup would return a price.
    px = _prices({DATES[0]: 100.0, DATES[1]: 125.0, DATES[2]: 126.0})
    check("holiday skipped", find_ko_day(accu, DATES, _Holidays([DATES[1]]), px), 2)


def isolated_render():
    print("Rendered KO row checks:")
    wd = _NeverHoliday()
    # Row 1 knocks out on DATES[5] (index 5 -> column T); row 2 never does.
    prices = {DATES[0]: 100.0, DATES[1]: 101.0, DATES[2]: 102.0, DATES[3]: 103.0,
              DATES[4]: 104.0, DATES[5]: 125.0, DATES[6]: 126.0, DATES[7]: 127.0,
              DATES[8]: 128.0, DATES[9]: 129.0}
    px = _prices(prices)
    records = [_rec(bank_doc='KO-ROW', ko=120.0), _rec(bank_doc='LIVE-ROW', ko=500.0)]

    wb = openpyxl.Workbook()
    ws = wb.active
    count_row = 3
    # Same call sequence build_workbook uses for a block.
    ko_offsets = _ko_offsets(records, DATES, wd, px)
    check("KO day located at the expected column", ko_offsets, {0: 5})
    _write_contracts(ws, records, 'ACCU', count_row, REPORT_DATE, DATES, wd,
                     price_lookup=px, bank_id='DBPe')
    _write_status_grid(ws, count_row, len(records), DATES, ko_offsets)

    ko_row = count_row + 4        # first data row
    live_row = ko_row + 1

    check("head count subtracts the KO'd row", ws[f'A{count_row}'].value,
          f'=2-COUNTIF(N{ko_row}:N{live_row},"*DONE*")-1')

    check("KO-day price kept", ws[f'T{ko_row}'].value, 125.0)
    check("KO-day price bold", ws[f'T{ko_row}'].font.b, True)
    check("next column holds the KO literal", ws[f'U{ko_row}'].value, 'KO')
    check("KO literal bold", ws[f'U{ko_row}'].font.b, True)
    check("KO literal not grey-filled", ws[f'U{ko_row}'].fill.patternType, None)

    after = [(c, ws[f'{c}{ko_row}'].value) for c in ('V', 'W', 'X')]
    check("days after the KO marker blank", after, [('V', None), ('W', None), ('X', None)])
    greys = [ws[f'{c}{ko_row}'].fill.fgColor.rgb for c in ('V', 'W', 'X')]
    check("days after the KO marker grey-filled", greys, [_GREY_FILL] * 3)

    colors = {c: ws[f'{c}{ko_row}'].font.color for c in _ALL_ROW_COLS}
    not_red = sorted(c for c, f in colors.items() if f is None or f.rgb != _RED)
    check("whole KO row A:AJ is red", not_red, [])

    live_colors = {c: ws[f'{c}{live_row}'].font.color for c in _ALL_ROW_COLS}
    reddened = sorted(c for c, f in live_colors.items() if f is not None and f.rgb == _RED)
    check("live row left black", reddened, [])
    check("live row keeps its last price", ws[f'X{live_row}'].value, 129.0)

    # A block with no knock-out keeps the unmodified head-count formula.
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    _write_contracts(ws2, [_rec(ko=500.0)], 'ACCU', count_row, REPORT_DATE, DATES, wd,
                     price_lookup=px, bank_id='DBPe')
    check("head count unchanged when nothing knocked out", ws2[f'A{count_row}'].value,
          f'=1-COUNTIF(N{ko_row}:N{ko_row},"*DONE*")')


def isolated_circles_stop_at_ko():
    print("Circle suppression after knock-out:")
    wd = _NeverHoliday()
    # Accumulator: strike 120, ko 125. Days 0-1 close under the strike (D),
    # day 2 breaches the barrier (KO), days 3-4 close under the strike again --
    # those must NOT be circled, the contract is dead by then.
    rec = _rec(spot=130.0, strike=120.0, ko=125.0)
    px = _prices({DATES[0]: 110.0, DATES[1]: 110.0, DATES[2]: 126.0,
                  DATES[3]: 110.0, DATES[4]: 110.0})
    cells = compute_status_flags([rec], 10, DATES, REPORT_DATE, wd, px)
    check("no D circles after the KO day", sorted(cells),
          sorted([('O', 10), ('P', 10)]))


def live_status_filter():
    print("Live contract_records status filter:")
    app = create_app()
    app.config["DATABASE"] = os.path.join(SERVER, "instance", "LTV Stocks.db")
    ctx = app.app_context()
    ctx.push()
    try:
        db = get_db()
        from ltv_app.blueprints.ltv_stocks.legacy_port.term_sheet_calc import contract_records

        bank_refs = [r[0] for r in db.execute(
            "SELECT ref_num FROM tbl_bank_account WHERE is_active = 1").fetchall()]
        leaked = set()
        for bank_ref in bank_refs:
            for product in ('ACCU', 'DECU'):
                for rec in contract_records(db, bank_ref, product):
                    if rec['status'] != 'active':
                        leaked.add(rec['status'])
        check("only active contracts reach the report", sorted(leaked), [])

        # The filter must actually be doing work -- if the DB holds no KO'd
        # contract at all, the check above passes vacuously and pins nothing.
        n_ko = db.execute(
            "SELECT COUNT(*) FROM tbl_stock_contract c "
            "INNER JOIN tbl_bank_account b ON c.bank_ref = b.ref_num "
            "WHERE c.status = 'KO' AND b.is_active = 1").fetchone()[0]
        check("DB holds KO'd contracts for the filter to exclude", n_ko > 0, True)
    finally:
        ctx.pop()


if __name__ == '__main__':
    isolated_find_ko_day()
    isolated_render()
    isolated_circles_stop_at_ko()
    live_status_filter()
    print("\nRESULT:", "PASS" if ok else "FAIL")
    sys.exit(0 if ok else 1)
