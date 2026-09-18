"""Verify the report generates both before and after the close is uploaded.

The report has to be usable at two moments in the day, and behave correctly at
each:

  * **Before the close.** No Yahoo Finance price has been stored for the report
    date yet, so the report date's O:X cell carries a live
    `=INDEX(closing_price!...)` lookup. Excel resolves it against the
    `closing_price` sheet when the file is opened, so the sheet is useful
    intraday.
  * **After the close.** `tbl_stock_price` has the day's close, so the cell
    carries that actual number and the sheet is a fixed record.

This is a requirement, not an accident, so it is pinned here. Two consequences
worth stating, because they look like defects and are not:

  * `find_ko_day` runs in Python against stored prices only. Before the close
    there is no price server-side, so the contract row cannot mark a knock-out
    on the report date -- while the Z:AJ grid, which Excel evaluates from the
    live formula, can. The grid knowing more than the row is the point of the
    grid, not a contradiction. After the close both read the same number.
  * USD contracts get no live formula (`closing_price` is quoted in the local
    currency), so their report-date cell is blank until the close is stored.

Runs entirely on synthetic records: no database, no client figures.

    .venv/Scripts/python.exe scripts/verify_report_date_price_cell.py
"""
import os
import sys
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from ltv_app.blueprints.ltv_stocks.legacy_port.excel_writer import _write_contracts
from ltv_app.tz import ph_today

ok = True


def check(label, actual, expected):
    global ok
    if actual == expected:
        print(f"  {label}: PASS")
    else:
        ok = False
        print(f"  {label}: FAIL  expected {expected!r}, got {actual!r}")


class _NeverHoliday:
    def is_holiday(self, d):
        return False


TODAY = ph_today()
# A ten-day window ending on the report date, which is today -- the state the
# report is in when it is generated during the session.
DATES = [TODAY - timedelta(days=9 - i) for i in range(10)]
COUNT_ROW = 3
ROW = COUNT_ROW + 4
LIVE_FORMULA = f'=INDEX(closing_price!A:C,MATCH(M{ROW},closing_price!A:A,),3)'


def _rec(**kw):
    base = dict(code_ref=1, stock_name='Test Stock NO GTD', stock_name_plain='Test Stock',
                code='0001', bank_doc='TEST-1', shares='100 / 200', spot=100.0,
                strike=80.0, ko=120.0, start_date=DATES[0], end_date=date(2027, 6, 1),
                received=1, total=12, yahoo_ticker='0001.HK', next_date=date(2026, 7, 1),
                frequency='monthly', ccy_id='HKD')
    base.update(kw)
    return base


def build(rec, prices, report_date=TODAY):
    wb = openpyxl.Workbook()
    ws = wb.active
    _write_contracts(ws, [rec], 'ACCU', COUNT_ROW, report_date, DATES, _NeverHoliday(),
                     price_lookup=lambda c, d: prices.get(d), bank_id='DBPe')
    return ws


PRICED_TO_YESTERDAY = {d: 100.0 for d in DATES[:9]}
PRICED_THROUGH_TODAY = {d: 100.0 for d in DATES}


def main():
    print("Before the close -- nothing stored for the report date:")
    ws = build(_rec(), PRICED_TO_YESTERDAY)
    check("report-date cell carries the live lookup", ws[f'X{ROW}'].value, LIVE_FORMULA)
    check("yesterday still shows its stored close", ws[f'W{ROW}'].value, 100.0)

    print("\nAfter the close -- Yahoo Finance price uploaded:")
    ws = build(_rec(), PRICED_THROUGH_TODAY)
    check("report-date cell carries the actual close", ws[f'X{ROW}'].value, 100.0)

    print("\nA past report date is a fixed record, never a live lookup:")
    past = DATES[8]
    ws = build(_rec(), PRICED_TO_YESTERDAY, report_date=past)
    check("stored close is used", ws[f'W{ROW}'].value, 100.0)
    ws = build(_rec(), {d: 100.0 for d in DATES[:8]}, report_date=past)
    check("an unpriced past date stays blank, not a formula",
          ws[f'W{ROW}'].value, None)

    print("\nUSD contracts take no live lookup:")
    ws = build(_rec(ccy_id='USD'), PRICED_TO_YESTERDAY)
    check("report-date cell left blank before the close", ws[f'X{ROW}'].value, None)
    ws = build(_rec(ccy_id='USD'), PRICED_THROUGH_TODAY)
    check("actual close used once stored", ws[f'X{ROW}'].value, 100.0)

    print("\nPASS" if ok else "\nFAIL")
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
