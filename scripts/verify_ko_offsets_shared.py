"""Verify the knock-out mapping is computed once per block and shared.

`build_workbook` computed `_ko_offsets` for a block to hand to
`_write_status_grid`, and `_write_contracts` then computed the identical dict
again internally. Two costs:

  * `find_ko_day` walks the ten-day window calling `price_lookup` per day, and
    that lookup is an uncached SELECT -- so every contract paid up to ten
    queries, twice.
  * More importantly it left two call sites that had to stay argument-identical.
    If they ever diverged, the red KO marking in the Z:AJ grid would drift out
    of step with the contract table's and nothing would notice.

So `_write_contracts` now accepts the caller's mapping. It still computes its
own when none is passed, which is what the other verify scripts rely on.

    .venv/Scripts/python.exe scripts/verify_ko_offsets_shared.py
"""
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from ltv_app import create_app
from ltv_app.blueprints.database.views import get_db
from ltv_app.blueprints.ltv_stocks.legacy_port import excel_writer
from ltv_app.blueprints.ltv_stocks.legacy_port.excel_writer import (
    _OX_COLS, _write_contracts, build_workbook)
from ltv_app.blueprints.ltv_stocks.legacy_port.term_sheet_calc import contract_records

SERVER = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BANK_ID = 'DBPe'
REPORT_DATE = date(2026, 7, 10)
ok = True


def check(label, actual, expected):
    global ok
    if actual == expected:
        print(f"  {label}: PASS")
    else:
        ok = False
        print(f"  {label}: FAIL  expected {expected!r}, got {actual!r}")


class _NeverHoliday:
    def is_holiday(self, d):
        return False


DATES = [date(2026, 6, 1), date(2026, 6, 2), date(2026, 6, 3), date(2026, 6, 4),
         date(2026, 6, 5), date(2026, 6, 8), date(2026, 6, 9), date(2026, 6, 10),
         date(2026, 6, 11), date(2026, 6, 12)]


def _rec(**kw):
    base = dict(code_ref=1, stock_name='Test Stock NO GTD', stock_name_plain='Test Stock',
                code='0001', bank_doc='TEST-1', shares='100 / 200', spot=100.0,
                strike=80.0, ko=120.0, start_date=DATES[0], end_date=date(2027, 6, 1),
                received=1, total=12, yahoo_ticker='0001.HK', next_date=date(2026, 7, 1),
                frequency='monthly', ccy_id='HKD')
    base.update(kw)
    return base


def honours_callers_mapping():
    """The caller's mapping wins -- that is what stops the two views drifting."""
    print("The caller's mapping is used, not recomputed:")
    prices = {d: 100.0 for d in DATES}          # nothing knocks out on its own
    ws = openpyxl.Workbook().active
    _write_contracts(ws, [_rec()], 'ACCU', 3, DATES[-1], DATES, _NeverHoliday(),
                     price_lookup=lambda c, d: prices.get(d), bank_id='DBPe',
                     ko_days={0: 2})
    check("marker written where the caller said, not where prices say",
          ws[f'{_OX_COLS[3]}7'].value, 'KO')

    ws = openpyxl.Workbook().active
    _write_contracts(ws, [_rec()], 'ACCU', 3, DATES[-1], DATES, _NeverHoliday(),
                     price_lookup=lambda c, d: prices.get(d), bank_id='DBPe')
    check("still computes its own when none is passed",
          any(ws[f'{c}7'].value == 'KO' for c in _OX_COLS), False)


def counted_once():
    print("\nfind_ko_day is called once per contract, not twice:")
    calls = []
    real = excel_writer.find_ko_day

    def counting(rec, *a, **kw):
        calls.append(rec['code_ref'])
        return real(rec, *a, **kw)

    app = create_app()
    app.config["DATABASE"] = os.path.join(SERVER, "instance", "LTV Stocks.db")
    ctx = app.app_context()
    ctx.push()
    try:
        db = get_db()
        bank_ref = db.execute(
            "SELECT ref_num FROM tbl_bank_account WHERE bank_id = ?", (BANK_ID,)).fetchone()[0]
        n = sum(len([r for r in contract_records(db, bank_ref, p) if r['ccy_id'] == ccy])
                for ccy in ('HKD', 'SGD') for p in ('ACCU', 'DECU'))
        excel_writer.find_ko_day = counting
        try:
            build_workbook(db, REPORT_DATE, [BANK_ID])
        finally:
            excel_writer.find_ko_day = real
    finally:
        ctx.pop()

    check(f"one find_ko_day per contract ({n} contracts)", len(calls), n)
    print(f"    (was {2 * n} before the fix)")


def main():
    honours_callers_mapping()
    counted_once()
    print("\nPASS" if ok else "\nFAIL")
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
