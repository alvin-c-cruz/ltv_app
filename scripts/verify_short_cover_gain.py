"""Regression harness: a short cover's gain is measured against the short it closes.

A cover (Buy (Pay Short)) realises the gain on shares sold short earlier, so
its basis is the short average carried just before it. The Trades Done Excel
and the printable Trades Done page used the day's closing average instead,
which differs only when a short sale follows the cover on the same day: on
2026-10-22 a cover closed an earlier short and a new short sale followed, and
the cover was measured against the new sale's price.

For every date with a cover, every cover row in both reports must carry, as
its average (Excel column O) and its cost / quantity (printable page), the
average from replaying the short book in ledger order (trade date, ref_num)
up to the row before it. Expected figures come from the database at runtime;
failures never print an amount -- the repository is public.

Run:   .venv/Scripts/python.exe scripts/verify_short_cover_gain.py [db_path]
       (default: instance/LTV Stocks.db)
"""
import os
import sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
SERVER = os.path.dirname(HERE)
sys.path.insert(0, SERVER)

from openpyxl import load_workbook

from ltv_app import create_app
from ltv_app.blueprints.database.views import get_db
from ltv_app.blueprints.transactions.extensions import TransactionSummary, TradesDoneReport
from ltv_app.blueprints.transactions.extensions import download_trades_done_with_gain_loss as excel

DB_PATH = os.path.join(SERVER, "instance", "LTV Stocks.db")
COVER = "Buy (Pay Short)"


def _open(db_path):
    app = create_app()
    app.config["DATABASE"] = db_path
    ctx = app.app_context()
    ctx.push()
    return ctx, get_db()


def _average_before(db, bank_ref, code_ref, trade_date, ref_num):
    # A short sale adds its net proceeds; a cover takes off its share of the
    # cost, leaving the average unchanged.
    rows = db.execute(
        "SELECT quantity, price, brokerage, commission, foreign_charge, stamp_duty, misc "
        "FROM tbl_transaction_short WHERE bank_ref = ? AND code_ref = ? "
        "AND (trade_date < ? OR (trade_date = ? AND ref_num < ?)) ORDER BY trade_date, ref_num",
        (bank_ref, code_ref, trade_date, trade_date, ref_num),
    ).fetchall()
    balance, cost, average = 0, 0.0, None
    for r in rows:
        qty = r["quantity"]
        charges = sum(r[k] or 0 for k in ("brokerage", "commission", "foreign_charge", "stamp_duty", "misc"))
        if qty < 0 and balance <= 0:
            cost += -qty * r["price"] - charges
        elif qty > 0 and balance < 0:
            cost = cost * (1 - qty / -balance) if -balance > qty else 0.0
        balance += qty
        if balance < 0:
            average = cost / -balance
    return average


def _expected(db, trade_date):
    """{(bank_name, stock, quantity, price): [average before each such cover]}, and
    whether a short sale follows a cover in the same stock that day."""
    rows = db.execute(
        "SELECT s.ref_num, s.bank_ref, s.code_ref, s.transaction_type, s.quantity, s.price, "
        "       b.bank_name, c.stock_name, c.code "
        "FROM tbl_transaction_short s "
        "INNER JOIN tbl_bank_account b ON b.ref_num = s.bank_ref "
        "INNER JOIN tbl_code c ON c.ref_num = s.code_ref "
        "WHERE s.trade_date = ? ORDER BY s.ref_num",
        (trade_date,),
    ).fetchall()
    expected, covered, sale_after_cover = defaultdict(list), set(), False
    for r in rows:
        key = (r["bank_ref"], r["code_ref"])
        if r["transaction_type"] == COVER:
            stock = f"{r['stock_name']} ({r['code']})"
            # Both reports carry the price at 4 dp (TransactionSummary formats it)
            expected[(r["bank_name"], stock, abs(r["quantity"]), round(r["price"], 4))].append(
                _average_before(db, r["bank_ref"], r["code_ref"], trade_date, r["ref_num"]))
            covered.add(key)
        elif r["transaction_type"] == "Sell (Short)" and key in covered:
            sale_after_cover = True
    return expected, sale_after_cover


def _same(got, want):
    def norm(d):
        return {k: sorted(round(v, 6) for v in vals if v is not None) for k, vals in d.items()}
    return norm(got) == norm(want)


def _excel_averages(db, trade_date, summary):
    excel.fetch_hkd_rate = lambda ccy: (1.0, "test")
    path = excel.DownloadTradesDoneWithGainLoss(db=db, trade_date=trade_date, trade_summary=summary).filename
    try:
        ws = load_workbook(path).active
    finally:
        os.remove(path)
    got = defaultdict(list)
    for r in range(1, ws.max_row + 1):
        if ws[f"D{r}"].value == COVER:
            got[(ws[f"A{r}"].value, ws[f"C{r}"].value, ws[f"G{r}"].value, ws[f"E{r}"].value)].append(ws[f"O{r}"].value)
    return got


def _report_averages(db, trade_date, summary):
    got = defaultdict(list)
    for block in TradesDoneReport(db, trade_date, summary).blocks:
        for bank in block.get("banks", []):
            for row in bank["rows"]:
                if row["transaction_type"] == COVER:
                    got[(row["bank_name"], row["stock_name"], row["quantity"], row["price"])].append(
                        row["cost"] / row["quantity"])
    return got


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    db_path = sys.argv[1] if len(sys.argv) > 1 else DB_PATH
    ctx, db = _open(db_path)
    real_feed = excel.fetch_hkd_rate
    try:
        dates = [r[0] for r in db.execute(
            "SELECT DISTINCT trade_date FROM tbl_transaction_short WHERE transaction_type = ? ORDER BY 1",
            (COVER,))]
        ok, covers, sale_after_cover = True, 0, []
        for d in dates:
            expected, after = _expected(db, d)
            if after:
                sale_after_cover.append(d)
            covers += sum(len(v) for v in expected.values())
            summary = TransactionSummary(db, d)
            if not _same(_excel_averages(db, d, summary), expected):
                print(f"  {d}: FAIL  Trades Done Excel: a cover's average is not the short it closes")
                ok = False
            if not _same(_report_averages(db, d, summary), expected):
                print(f"  {d}: FAIL  printable Trades Done: a cover's cost is not the short it closes")
                ok = False
        if not dates:
            print("  FAIL: no date with a short cover")
            ok = False
        print(f"  {len(dates)} cover date(s), {covers} cover(s) checked in both reports")
        print(f"  short sale after a same-day cover: {', '.join(sale_after_cover) or 'none in this database'}")
    finally:
        excel.fetch_hkd_rate = real_feed
        ctx.pop()

    print("RESULT:", "ALL PASS" if ok else "FAIL")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
