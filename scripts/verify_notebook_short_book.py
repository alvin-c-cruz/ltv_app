"""Regression harness: short-book trades must print in the notebook.

Short sales and their covers are kept in tbl_transaction_short (the Short
Book). The notebook read only tbl_transaction, so a Sell (Short) or a Buy (Pay
Short) never printed and had to be typed in by hand. Found 2026-10-22 on the
2026-10-21 notebook: that day's pay short was missing.

For every date with a Sell (Short) or Buy (Pay Short), each such row must
print under its own account with Alvin's label, its quantity and its price,
and the (account, stock) short block must open at the short balance before
the date and end at the balance after that day's short-book trades. After a
short sale that leaves a short, the block carries "Short / Average / <short
average>" in L, as Alvin added it by hand on the 22 Oct 2026 notebook; any
other block has none. Expected figures are derived from the database at
runtime; failures never print a balance -- the repository is public.

Each notebook is generated into instance/temp/ exactly as the /notebook/ route
does; a file this script created is removed after it is read. The database is
only read.

Run:   .venv/Scripts/python.exe scripts/verify_notebook_short_book.py [db_path] [--all]
       (default: instance/LTV Stocks.db, the 25 most recent short-book dates)
"""
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SERVER = os.path.dirname(HERE)
sys.path.insert(0, SERVER)

from openpyxl import load_workbook

from ltv_app import create_app
from ltv_app.blueprints.database.views import get_db
from ltv_app.blueprints.notebook.extensions import CreateNotebook, get_transactions, get_transfers

DB_PATH = os.path.join(SERVER, "instance", "LTV Stocks.db")
RECENT = 25

# The labels Alvin wrote by hand (13 Oct and 10 Feb 2026 notebooks).
LABELS = {"Sell (Short)": "Sell Short (Spot)", "Buy (Pay Short)": "Pay Short (Spot)"}


def _open(db_path):
    app = create_app()
    app.config["DATABASE"] = db_path
    ctx = app.app_context()
    ctx.push()
    return ctx, get_db()


def _short_dates(db):
    rows = db.execute(
        "SELECT DISTINCT trade_date FROM tbl_transaction_short "
        "WHERE transaction_type IN (?, ?) ORDER BY trade_date",
        tuple(LABELS),
    ).fetchall()
    return [str(r[0])[:10] for r in rows]


def _short_average(db, bank_name, code, trade_date):
    # Replays the short book in ledger order: a short sale adds its net proceeds,
    # a cover takes off its share of the cost, leaving the average unchanged.
    rows = db.execute(
        "SELECT s.quantity, s.price, s.brokerage, s.commission, s.foreign_charge, s.stamp_duty, s.misc "
        "FROM tbl_transaction_short s "
        "INNER JOIN tbl_bank_account b ON b.ref_num = s.bank_ref "
        "INNER JOIN tbl_code c ON c.ref_num = s.code_ref "
        "WHERE b.bank_name = ? AND c.code = ? AND s.trade_date <= ? ORDER BY s.trade_date, s.ref_num",
        (bank_name, code, trade_date),
    ).fetchall()
    balance, cost, average = 0, 0.0, None
    for r in rows:
        qty = r["quantity"]
        charges = sum(r[k] or 0 for k in ("brokerage", "commission", "foreign_charge", "stamp_duty", "misc"))
        if qty < 0 and balance <= 0:
            cost += -qty * r["price"] - charges
        elif qty > 0 and balance < 0:
            cost = cost * (1 - qty / -balance) if -balance > qty else 0.0
        balance += qty
        if balance < 0:
            average = cost / -balance
    return average


def _expected(db, trade_date):
    rows = db.execute(
        "SELECT s.ref_num, b.bank_name, cy.ccy_id, c.code, c.stock_name, "
        "       s.transaction_type, s.quantity, s.price "
        "FROM tbl_transaction_short s "
        "INNER JOIN tbl_bank_account b ON b.ref_num = s.bank_ref "
        "INNER JOIN tbl_code c ON c.ref_num = s.code_ref "
        "INNER JOIN tbl_currency cy ON cy.ref_num = c.ccy_ref "
        "WHERE s.trade_date = ? AND s.transaction_type IN (?, ?) "
        "ORDER BY s.ref_num",
        (trade_date,) + tuple(LABELS),
    ).fetchall()

    blocks = {}
    for r in rows:
        key = (r["bank_name"], r["code"])
        if key not in blocks:
            before = db.execute(
                "SELECT SUM(s.quantity) FROM tbl_transaction_short s "
                "INNER JOIN tbl_bank_account b ON b.ref_num = s.bank_ref "
                "INNER JOIN tbl_code c ON c.ref_num = s.code_ref "
                "WHERE b.bank_name = ? AND c.code = ? AND s.trade_date < ?",
                (r["bank_name"], r["code"], trade_date),
            ).fetchone()[0]
            before = int(before) if before is not None else 0
            header = r["bank_name"] if r["ccy_id"] == "HKD" else f"{r['bank_name']} ({r['ccy_id']})"
            blocks[key] = {"header": header, "before": before, "after": before, "rows": [],
                           "short_sale": False}
        blocks[key]["after"] += int(r["quantity"])
        blocks[key]["rows"].append(r)
        blocks[key]["short_sale"] |= r["transaction_type"] == "Sell (Short)"
    for (bank_name, code), block in blocks.items():
        block["average"] = (_short_average(db, bank_name, code, trade_date)
                            if block["short_sale"] and block["after"] < 0 else None)
    return blocks


def _printed(path):
    # Every titled line, with the account header it sits under, and every
    # remaining-shares line (F holds the "shares"/"short" formula).
    ws = load_workbook(path)["notebook"]
    titles, results = [], []
    header = None
    for r in range(1, ws.max_row + 1):
        a, c, f = ws[f"A{r}"].value, ws[f"C{r}"].value, ws[f"F{r}"].value
        if isinstance(a, str) and r > 1:
            header = a
        if ws[f"B{r}"].value is not None and isinstance(c, str):
            titles.append((r, header, c, ws[f"R{r + 1}"].value, ws[f"S{r + 1}"].value))
        if isinstance(f, str) and f.startswith('=IF(G') and '"short"' in f:
            results.append(r)
    return ws, titles, results


def _evaluate(ws, formula):
    # "=G132-G133+G134": the first term is the opening figure, each later term
    # is "=R<n>", the quantity of one printed trade.
    terms = re.findall(r"([=+-])G(\d+)", formula)
    opening = ws[f"G{terms[0][1]}"].value
    total = opening if isinstance(opening, (int, float)) else None
    for sign, row in terms[1:]:
        ref = ws[f"G{row}"].value
        qty = ws[ref[1:]].value if isinstance(ref, str) and ref.startswith("=R") else None
        if total is None or qty is None:
            return opening, None
        total += qty if sign == "+" else -qty
    return opening, total


def _check_date(db, trade_date):
    blocks = _expected(db, trade_date)
    path = None
    try:
        transactions = get_transactions(db=db, trade_date=trade_date)
        transfers = get_transfers(db=db, trade_date=trade_date)
        path = CreateNotebook(db=db, trade_date=trade_date,
                              transactions=transactions, transfers=transfers).filename
        ws, titles, results = _printed(path)

        ok = True
        averaged = []
        for (bank_name, code), block in blocks.items():
            last = None
            for row in block["rows"]:
                want = f"{LABELS[row['transaction_type']]} {row['stock_name']} ({code})"
                hit = next((t for t in titles
                            if t[1] == block["header"] and t[2] == want
                            and t[3] == abs(row["quantity"]) and t[4] == row["price"]), None)
                if hit is None:
                    print(f"  {trade_date} short #{row['ref_num']} {code}: FAIL  not printed under its account")
                    ok = False
                    continue
                titles.remove(hit)
                last = max(last or 0, hit[0])
            if last is None:
                continue
            result_row = next((r for r in results if r > last), None)
            if result_row is None:
                print(f"  {trade_date} {code}: FAIL  no balance block after the short trades")
                ok = False
                continue
            opening, closing = _evaluate(ws, ws[f"G{result_row}"].value)
            if opening != block["before"]:
                print(f"  {trade_date} {code}: FAIL  block does not open at the short balance before the date")
                ok = False
            if closing != block["after"]:
                print(f"  {trade_date} {code}: FAIL  block does not end at the short balance after the date")
                ok = False
            printed = [ws[f"L{result_row + i}"].value for i in (-2, -1, 0)]
            if block["average"] is None:
                if any(v is not None for v in printed):
                    print(f"  {trade_date} {code}: FAIL  short average printed without a short sale leaving a short")
                    ok = False
            elif (printed[:2] != ["Short", "Average"] or not isinstance(printed[2], (int, float))
                  or abs(printed[2] - block["average"]) > 1e-6):
                print(f"  {trade_date} {code}: FAIL  short average missing or not the ledger's")
                ok = False
            else:
                averaged.append(trade_date)
    finally:
        if path and os.path.exists(path):
            os.remove(path)
    return ok, sum(len(b["rows"]) for b in blocks.values()), len(averaged)


def main():
    args = [a for a in sys.argv[1:] if a != "--all"]
    db_path = args[0] if args else DB_PATH
    ctx, db = _open(db_path)
    try:
        dates = _short_dates(db)
        if not dates:
            print("  FAIL: no date with a short-book trade")
            sys.exit(1)
        if "--all" not in sys.argv:
            dates = dates[-RECENT:]

        ok, rows, averages = True, 0, 0
        for d in dates:
            date_ok, n, a = _check_date(db, d)
            ok &= date_ok
            rows += n
            averages += a
        print(f"  {len(dates)} short-book date(s), {rows} short trade(s) checked, {averages} short average(s)")
    finally:
        ctx.pop()

    print("RESULT:", "ALL PASS" if ok else "FAIL")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
