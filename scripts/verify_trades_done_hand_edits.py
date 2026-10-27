"""Regression harness: the Trades Done with gain loss Excel comes out the way
Alvin finishes it by hand, so the issued file needs no edits.

Alvin, 2026-10-27, after the 26 Oct comparison of his saved files against raw
downloads (BUGS.md 2026-10-26), asked for all of these:

  E  E2 is a fixed time, not =NOW() (which changes whenever the file is
     opened): the moment the file is made when it is made on the trade date,
     the 4:00 PM close otherwise. Format h:mm AM/PM.
  F  A gain/loss section with ONE bank subtotal leaves Cost / Gain / % off its
     TOTAL footer -- the subtotal just above already shows them. With two or
     more bank subtotals the footer keeps them. G and H stay either way.
  S  A section with ONE trade line leaves Cost / Gain / % off its bank subtotal
     row, which would only repeat that line; the black divider stays. With two
     or more trade lines every bank subtotal keeps them.
  H  The "Transfer of Stocks" heading is left-aligned, not centred.

F and S sweep every gain/loss section (Sell (Spot), and Buy (Pay Short) on the
short book) of every 2026 trade date, and each also requires both of its shapes
to occur, so neither can pass by having nothing to look at. The FX feed is
replaced by a fixed test rate, so nothing is fetched. Workbooks are generated
through the route, so this does not depend on where the file is written.
Failures print no figure -- the repository is public.

Run: server/.venv/Scripts/python.exe scripts/verify_trades_done_hand_edits.py
"""
import io
import os
import sys
from datetime import datetime, time

HERE = os.path.dirname(os.path.abspath(__file__))
SERVER = os.path.dirname(HERE)
sys.path.insert(0, SERVER)

from openpyxl import load_workbook

from ltv_app import create_app
from ltv_app.blueprints.transactions.extensions import download_trades_done_with_gain_loss as module
from ltv_app.tz import PH_TZ

DB_PATH = os.path.join(SERVER, "instance", "LTV Stocks.db")
SECTION_LABELS = ("Cost ", "Short Sale ")
GAIN_COLS = ("L", "M", "N")

app = create_app({"DATABASE": DB_PATH, "LOGIN_DISABLED": True})
client = app.test_client()
module.fetch_hkd_rate = lambda ccy: (6.1234, "test feed")
ok = True


def _case(label, actual, expected):
    global ok
    if actual == expected:
        print(f"  {label}: PASS")
    else:
        ok = False
        print(f"  {label}: FAIL")


def _sheet(trade_date):
    # Removes the copy the route leaves in instance/temp/, unless it was already there.
    leftover = os.path.join(app.instance_path, "temp", f"{trade_date} Trades Done with gain loss.xlsx")
    existed = os.path.exists(leftover)
    r = client.get(f"/trades/download_with_gain_loss/{trade_date}")
    data = r.get_data()
    r.close()
    if not existed and os.path.exists(leftover):
        os.remove(leftover)
    if r.status_code != 200:
        return None
    return load_workbook(io.BytesIO(data)).worksheets[0]


def _dates(sql):
    with app.app_context():
        from ltv_app.blueprints.database.views import get_db
        return [r[0] for r in get_db().execute(sql).fetchall()]


def _empty(cell):
    b = cell.border
    return cell.value is None and not any(s is not None and s.style for s in (b.left, b.right, b.top, b.bottom))


def _is_divider(ws, row):
    fill = ws[f"A{row}"].fill
    return fill is not None and fill.patternType == "solid" and str(fill.fgColor.rgb).endswith("000000")


def _sections(ws):
    """Each gain/loss section as (trade-line rows, bank-subtotal rows, footer row or None)."""
    out = []
    for row in range(1, ws.max_row + 1):
        label = ws[f"L{row}"].value
        if not (isinstance(label, str) and label.startswith(SECTION_LABELS)):
            continue
        lines, subtotals, footer = [], [], None
        r = row + 1
        while True:
            if ws[f"E{r}"].value == "TOTAL":
                footer = r
                break
            if _is_divider(ws, r):
                subtotals.append(r)
            elif ws[f"D{r}"].value:
                lines.append(r)
            else:
                break
            r += 1
        out.append((lines, subtotals, footer))
    return out


gain_dates = _dates(
    "SELECT DISTINCT substr(trade_date, 1, 10) FROM ("
    " SELECT trade_date FROM tbl_transaction WHERE transaction_type = 'Sell (Spot)'"
    " UNION ALL SELECT trade_date FROM tbl_transaction_short WHERE transaction_type = 'Buy (Pay Short)')"
    " WHERE trade_date >= '2026-01-01' ORDER BY 1")
transfer_dates = _dates(
    "SELECT DISTINCT substr(trade_date, 1, 10) FROM tbl_transaction "
    "WHERE transaction_type LIKE 'Transfer%' AND trade_date >= '2026-01-01' ORDER BY 1")

# --- E: the time in E2.
print("E2 time:")
ws = _sheet(gain_dates[-1]) if gain_dates else None
_case("E0 a trade date exists to test", ws is not None, True)
if ws is not None:
    _case("E1 a past trade date shows the 4:00 PM close, as a value",
          ws["E2"].value, time(16, 0))
    _case("E2 formatted h:mm AM/PM", ws["E2"].number_format, "h:mm AM/PM")
    real_now = getattr(module, "ph_now", None)
    made = datetime.fromisoformat(gain_dates[-1]).replace(hour=13, minute=43, second=27, tzinfo=PH_TZ)
    module.ph_now = lambda: made
    try:
        ws = _sheet(gain_dates[-1])
    finally:
        if real_now is None:
            del module.ph_now
        else:
            module.ph_now = real_now
    _case("E3 made on the trade date, it shows the minute it was made",
          ws["E2"].value if ws is not None else None, time(13, 43))

# --- F and S: every gain/loss section of 2026.
print("Gain/loss sections, every 2026 trade date:")
shapes = {"one subtotal": 0, "several subtotals": 0, "one line": 0, "several lines": 0}
wrong = {"F1": 0, "F2": 0, "F3": 0, "S1": 0, "S2": 0}
for trade_date in gain_dates:
    ws = _sheet(trade_date)
    if ws is None:
        continue
    for lines, subtotals, footer in _sections(ws):
        if footer is not None:
            if not all(str(ws[f"{c}{footer}"].value or "").startswith("=SUM(") for c in ("G", "H")):
                wrong["F3"] += 1
            if len(subtotals) == 1:
                shapes["one subtotal"] += 1
                if not all(_empty(ws[f"{c}{footer}"]) for c in GAIN_COLS):
                    wrong["F1"] += 1
            elif len(subtotals) > 1:
                shapes["several subtotals"] += 1
                if not all(str(ws[f"{c}{footer}"].value or "").startswith("=") for c in GAIN_COLS):
                    wrong["F2"] += 1
        if len(lines) == 1:
            shapes["one line"] += 1
            if not all(_empty(ws[f"{c}{r}"]) and _is_divider(ws, r) for c in GAIN_COLS for r in subtotals):
                wrong["S1"] += 1
        elif len(lines) > 1:
            shapes["several lines"] += 1
            if not all(str(ws[f"{c}{r}"].value or "").startswith("=") for c in GAIN_COLS for r in subtotals):
                wrong["S2"] += 1
for shape, count in shapes.items():
    _case(f"  a section with {shape} exists to test", count > 0, True)
_case("F1 one bank subtotal: footers still carrying Cost / Gain / %", wrong["F1"], 0)
_case("F2 several bank subtotals: footers missing Cost / Gain / %", wrong["F2"], 0)
_case("F3 footers missing the shares or total sum", wrong["F3"], 0)
_case("S1 one trade line: subtotal rows still carrying Cost / Gain / %", wrong["S1"], 0)
_case("S2 several trade lines: subtotal rows missing Cost / Gain / %", wrong["S2"], 0)

# --- H: the transfer heading.
print("Transfer heading:")
ws = _sheet(transfer_dates[-1]) if transfer_dates else None
heading = None
if ws is not None:
    heading = next((ws[f"A{r}"] for r in range(1, ws.max_row + 1)
                    if ws[f"A{r}"].value == "Transfer of Stocks"), None)
_case("H0 a report with transfers exists to test", heading is not None, True)
if heading is not None:
    _case("H1 'Transfer of Stocks' is left-aligned", heading.alignment.horizontal, "left")

print("RESULT:", "ALL PASS" if ok else "FAIL")
sys.exit(0 if ok else 1)
