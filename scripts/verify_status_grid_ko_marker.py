"""Verify the Z:AJ status grid does not read the KO marker as a price.

`_write_contracts` writes the text literal "KO" into the O:X price cell for the
day after a knock-out. `_write_status_grid` pairs each AA:AJ status cell to that
same O:X cell and reads it as a price.

Excel ranks text above every number, so an unguarded formula misreads it. On a
decumulator the branch `IF($G>=price,"KO",IF($F<=price,"D","."))` evaluates
`90>="KO"` as FALSE and `130<="KO"` as TRUE, printing **D** -- "strike breached,
still alive" -- for a contract that has already terminated. The accumulator
branch prints "KO", but only because `$G<="KO"` is TRUE for any number, i.e.
right by accident.

The `=0` and `=""` guards do not catch text, so the formula needs an explicit
ISTEXT guard.

Runs entirely on synthetic records: no database, no client figures.

    .venv/Scripts/python.exe scripts/verify_status_grid_ko_marker.py
"""
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from ltv_app.blueprints.ltv_stocks.legacy_port.excel_writer import (
    _OX_COLS, _STATUS_COLS, _ko_offsets, _write_contracts, _write_status_grid)

ok = True


def check(label, actual, expected):
    global ok
    if actual == expected:
        print(f"  {label}: PASS")
    else:
        ok = False
        print(f"  {label}: FAIL  expected {expected!r}, got {actual!r}")


class _NeverHoliday:
    def is_holiday(self, d):
        return False


DATES = [date(2026, 6, 1), date(2026, 6, 2), date(2026, 6, 3), date(2026, 6, 4),
         date(2026, 6, 5), date(2026, 6, 8), date(2026, 6, 9), date(2026, 6, 10),
         date(2026, 6, 11), date(2026, 6, 12)]
COUNT_ROW = 3


def _rec(**kw):
    base = dict(code_ref=1, stock_name='Test Stock NO GTD', stock_name_plain='Test Stock',
                code='0001', bank_doc='TEST-1', shares='100 / 200', spot=100.0,
                strike=80.0, ko=120.0, start_date=DATES[0], end_date=date(2027, 6, 1),
                received=1, total=12, yahoo_ticker='0001.HK', next_date=date(2026, 7, 1),
                frequency='monthly', ccy_id='HKD')
    base.update(kw)
    return base


def build(rec, product, prices):
    wb = openpyxl.Workbook()
    ws = wb.active
    lookup = lambda c, d: prices.get(d)
    ko = _ko_offsets([rec], DATES, _NeverHoliday(), lookup)
    _write_contracts(ws, [rec], product, COUNT_ROW, DATES[-1], DATES,
                     _NeverHoliday(), price_lookup=lookup, bank_id='DBPe')
    _write_status_grid(ws, COUNT_ROW, 1, DATES, ko_offsets=ko)
    return ws, ko.get(0)


# Accumulator: spot > strike, knocks out when the close reaches ko=120.
ACCU = (_rec(spot=100.0, strike=80.0, ko=120.0),
        'ACCU',
        {DATES[0]: 100.0, DATES[1]: 110.0, DATES[2]: 125.0})
# Decumulator: spot < strike, knocks out when the close falls to ko=90.
DECU = (_rec(spot=100.0, strike=130.0, ko=90.0),
        'DECU',
        {DATES[0]: 100.0, DATES[1]: 98.0, DATES[2]: 85.0})


def main():
    r = COUNT_ROW + 4

    for rec, product, prices in (ACCU, DECU):
        print(f"{product} block:")
        ws, ko_idx = build(rec, product, prices)
        check("knock-out located", ko_idx is not None, True)

        marker_col = _OX_COLS[ko_idx + 1]
        check("price cell after the KO day holds the text marker",
              ws[f'{marker_col}{r}'].value, 'KO')

        # Every status cell must refuse a text price, not just the marker one:
        # the guard is what makes the grid safe, wherever the text lands.
        missing = [c for c, p in zip(_STATUS_COLS, _OX_COLS)
                   if not str(ws[f'{c}{r}'].value).startswith(f'=IF(ISTEXT({p}{r}),"",')]
        check("every status cell guards against a text price", missing, [])

        f = str(ws[f'{_STATUS_COLS[ko_idx + 1]}{r}'].value)
        check("marker column's status formula is balanced",
              f.count('('), f.count(')'))
        print()

    print("PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
