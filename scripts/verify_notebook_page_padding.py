"""Regression harness: the notebook's ruled lines and print area run to the bottom
of its last printed page.

Alvin prints the notebook on 5.5 x 8.5 paper and, until 2026-10-27, carried
border_line()'s thin grey rules and the 17.25 row height down to the end of the
last page by hand, and set the print area to match. Across his October closing
files the print area ends on row 73, 105, 137 or 169 in 15 of 17: the first two
pages hold rows 1-73 (rows 5-23 are the 12 pt stock card), and every later page
32 more rows under the repeated title rows 1:2. He chose that rule on
2026-10-27 (BUGS.md 2026-10-26, item 1).

For each of the newest trade dates with a notebook, generated through the route:
  P1  the print area ends on the first page end (73, 105, 137 ...) at or below
      the last row with content;
  P2  every row after the content down to that end is ruled (top and bottom
      border on A:Q) and 17.25 high;
  P3  nothing is ruled below it.
Both a notebook that ends within the first 73 rows and a longer one are required,
so neither shape can pass by its absence. No figure is printed.

Run: server/.venv/Scripts/python.exe scripts/verify_notebook_page_padding.py
"""
import io
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SERVER = os.path.dirname(HERE)
sys.path.insert(0, SERVER)

from openpyxl import load_workbook

from ltv_app import create_app

DB_PATH = os.path.join(SERVER, "instance", "LTV Stocks.db")
FIRST_PAGES_END = 73
PAGE_ROWS = 32
ROW_HEIGHT = 17.25
COLUMNS = "ABCDEFGHIJKLMNOPQ"
DATES_TO_CHECK = 30

app = create_app({"DATABASE": DB_PATH, "LOGIN_DISABLED": True})
client = app.test_client()
ok = True


def _case(label, actual, expected):
    global ok
    if actual == expected:
        print(f"  {label}: PASS")
    else:
        ok = False
        print(f"  {label}: FAIL")


def _notebook(trade_date):
    # Removes the copy the route leaves in instance/temp/, unless it was already there.
    leftover = os.path.join(app.instance_path, "temp", f"{trade_date} notebook.xlsx")
    existed = os.path.exists(leftover)
    r = client.get(f"/notebook/generate/{trade_date}")
    data = r.get_data()
    r.close()
    if not existed and os.path.exists(leftover):
        os.remove(leftover)
    if r.status_code != 200:
        return None
    return load_workbook(io.BytesIO(data))["notebook"]


def _page_end(row):
    if row <= FIRST_PAGES_END:
        return FIRST_PAGES_END
    return FIRST_PAGES_END + PAGE_ROWS * math.ceil((row - FIRST_PAGES_END) / PAGE_ROWS)


def _ruled(ws, row):
    for col in COLUMNS:
        b = ws[f"{col}{row}"].border
        if not (b.top is not None and b.top.style and b.bottom is not None and b.bottom.style):
            return False
    return True


def _last_content_row(ws):
    last = max((r for r in range(1, ws.max_row + 1)
                if any(ws[f"{c}{r}"].value is not None for c in COLUMNS)), default=0)
    for rng in ws.merged_cells.ranges:
        last = max(last, rng.max_row)
    return last


with app.app_context():
    from ltv_app.blueprints.database.views import get_db
    dates = [r[0] for r in get_db().execute(
        "SELECT DISTINCT substr(trade_date, 1, 10) FROM tbl_transaction "
        "ORDER BY 1 DESC").fetchall()]

checked = {"short": 0, "long": 0}
wrong = {"P1": 0, "P2": 0, "P3": 0}
for trade_date in dates:
    if sum(checked.values()) >= DATES_TO_CHECK:
        break
    ws = _notebook(trade_date)
    if ws is None:
        continue
    last = _last_content_row(ws)
    end = _page_end(last)
    checked["short" if last <= FIRST_PAGES_END else "long"] += 1
    area = str(ws.print_area or "")
    if not area.endswith(f"$Q${end}"):
        wrong["P1"] += 1
    if not all(_ruled(ws, r) and ws.row_dimensions[r].height == ROW_HEIGHT
               for r in range(last + 1, end + 1)):
        wrong["P2"] += 1
    if _ruled(ws, end + 1):
        wrong["P3"] += 1

_case("P0a a notebook ending within the first 73 rows exists to test", checked["short"] > 0, True)
_case("P0b a longer notebook exists to test", checked["long"] > 0, True)
_case("P1 notebooks whose print area does not end on the page end", wrong["P1"], 0)
_case("P2 notebooks with an unruled or short row before the page end", wrong["P2"], 0)
_case("P3 notebooks ruled past the page end", wrong["P3"], 0)

print("RESULT:", "ALL PASS" if ok else "FAIL")
sys.exit(0 if ok else 1)
