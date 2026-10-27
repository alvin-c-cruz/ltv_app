"""Regression harness: the Trades Done Excel must total other currencies in HKD.

Alvin, 2026-10-22: "when there's other currency. there should be total in HKD.
put a currency converter with today's exchange rate so user can change it
later". His hand-made 22 Oct file showed the layout: SGD after all the HKD
sections, an HKD Sub Total, the SGD cost and gain/loss times a rate, and an
HKD and SGD Grand Total. The rate now sits in its own cell, filled from
today's feed (fetch_hkd_rate) and typed over by the user.

Checks, on every date in the database with a gain/loss section in a currency
other than HKD (the feed is replaced by a fixed test rate, so nothing is
fetched):
  1. Sections are grouped by currency, HKD first.
  2. HKD Sub Total = the sum of the HKD gain/loss rows' Cost and Gain / Loss.
  3. Each <CCY> -> HKD row = that currency's rows times the rate cell, and
     the rate cell holds the feed's rate.
  4. The Grand Total = the HKD Sub Total plus every converted row.
  5. With the feed down, the rate cell reads "enter rate" and every total
     that depends on it shows an error instead of a plausible figure.
Plus, on every other date: no total rows are written and no rate is fetched;
and fetch_hkd_rate returns None on a network error or an unusable reply.

Formulas are evaluated here (Excel's SUBTOTAL skips cells holding another
SUBTOTAL), as openpyxl does not calculate. Failures never print an amount --
the repository is public.

Run:   .venv/Scripts/python.exe scripts/verify_trades_done_hkd_total.py [db_path]
       (default: instance/LTV Stocks.db)
"""
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SERVER = os.path.dirname(HERE)
sys.path.insert(0, SERVER)

from openpyxl import load_workbook

from ltv_app import create_app
from ltv_app.blueprints.database.views import get_db
from ltv_app.blueprints.transactions.extensions import TransactionSummary
from ltv_app.blueprints.transactions.extensions import download_trades_done_with_gain_loss as module

DB_PATH = os.path.join(SERVER, "instance", "LTV Stocks.db")
TEST_RATE = 6.1234
GAIN_TYPES = ("Sell (Spot)", "Buy (Pay Short)")


class Error(Exception):
    pass


class Sheet:
    """Values of a generated sheet, with the few formula shapes it writes."""

    def __init__(self, ws):
        self.ws = ws
        self.cache = {}

    def value(self, ref):
        ref = ref.replace("$", "")
        if ref not in self.cache:
            raw = self.ws[ref].value
            self.cache[ref] = self.evaluate(raw[1:]) if isinstance(raw, str) and raw.startswith("=") else raw
        return self.cache[ref]

    def subtotal(self, col, first, last):
        total = 0
        for row in range(first, last + 1):
            raw = self.ws[f"{col}{row}"].value
            if isinstance(raw, str) and "SUBTOTAL(" in raw:
                continue
            value = self.value(f"{col}{row}")
            if isinstance(value, (int, float)):
                total += value
        return total

    def evaluate(self, formula):
        if formula.startswith(("IF(", '"', "TEXT(")) or "&" in formula:
            return formula
        expr = re.sub(r"SUBTOTAL\(9,([A-Z])(\d+):[A-Z](\d+)\)", r'_sub("\1",\2,\3)', formula)
        expr = re.sub(r"(?<![\"A-Za-z_])\$?([A-Z])\$?(\d+)", r'_val("\1\2")', expr)
        try:
            result = eval(expr, {"_sub": self.subtotal, "_val": self.value})
        except (TypeError, ZeroDivisionError):
            raise Error("#VALUE!")
        # Python repeats a string times an int (0 * "enter rate" == ""); Excel
        # gives #VALUE! for any arithmetic on text.
        if isinstance(result, str):
            raise Error("#VALUE!")
        return result

    def safe(self, ref):
        try:
            return self.value(ref)
        except Error as e:
            return str(e)


def _open(db_path):
    app = create_app()
    app.config["DATABASE"] = db_path
    ctx = app.app_context()
    ctx.push()
    return ctx, get_db()


def _dates(db):
    rows = db.execute(
        "SELECT DISTINCT substr(trade_date,1,10) FROM ("
        " SELECT trade_date FROM tbl_transaction UNION ALL SELECT trade_date FROM tbl_transaction_short"
        " UNION ALL SELECT trade_date FROM tbl_stock_contract) ORDER BY 1").fetchall()
    return [r[0] for r in rows]


def _foreign_gain(summary):
    long_book, short_book = summary.transactions, summary.short_transactions
    return sorted(ccy for ccy in set(long_book) | set(short_book) if ccy != "HKD" and (
        "Sell (Spot)" in long_book.get(ccy, {}) or "Buy (Pay Short)" in short_book.get(ccy, {})))


def _generate(db, trade_date, summary, rate):
    calls = []

    def feed(ccy):
        calls.append(ccy)
        return None if rate is None else (rate, "test feed")

    module.fetch_hkd_rate = feed
    ws = load_workbook(module.DownloadTradesDoneWithGainLoss(
        db=db, trade_date=trade_date, trade_summary=summary).output).active
    return ws, calls


def _layout(ws):
    """Section headers (row, ccy), gain/loss detail rows (row, ccy), and the
    labelled total rows {label: row}."""
    sections, details, labels = [], [], {}
    ccy = None
    for r in range(1, ws.max_row + 1):
        a, d = ws[f"A{r}"].value, ws[f"D{r}"].value
        if a in ("BUY", "SELL"):
            ccy = d
            sections.append((r, ccy))
        elif d in GAIN_TYPES and isinstance(ws[f"L{r}"].value, str):
            details.append((r, ccy))
        for col in ("J", "K"):
            v = ws[f"{col}{r}"].value
            if isinstance(v, str) and ("Total" in v or "→ HKD" in v):
                labels[v] = r
    return sections, details, labels


def _close(a, b):
    return isinstance(a, (int, float)) and abs(a - b) < 0.005


def _check(db, trade_date, summary, foreign):
    ok = True

    def fail(msg):
        nonlocal ok
        ok = False
        print(f"  {trade_date}: FAIL  {msg}")

    ws, calls = _generate(db, trade_date, summary, TEST_RATE)
    sheet = Sheet(ws)
    sections, details, labels = _layout(ws)

    ccys = [c for _, c in sections]
    grouped = list(dict.fromkeys(ccys))
    if ccys != sorted(ccys, key=grouped.index) or (ccys and "HKD" in ccys and ccys[0] != "HKD"):
        fail("sections not grouped by currency, HKD first")
    if sorted(calls) != foreign:
        fail(f"rate fetched for {calls}, expected {foreign}")

    def sums(ccy):
        rows = [r for r, c in details if c == ccy]
        return (sum(sheet.value(f"L{r}") for r in rows), sum(sheet.value(f"M{r}") for r in rows))

    totals = []
    if any(c == "HKD" for _, c in details):
        row = labels.get("HKD Sub Total")
        if row is None:
            fail("no HKD Sub Total")
        else:
            cost, gain = sums("HKD")
            if not (_close(sheet.value(f"L{row}"), cost) and _close(sheet.value(f"M{row}"), gain)):
                fail("HKD Sub Total is not the sum of the HKD gain/loss rows")
            totals.append(("HKD", row, cost, gain))

    for ccy in foreign:
        row = labels.get(f"{ccy} → HKD")
        if row is None:
            fail(f"no {ccy} → HKD row")
            continue
        if ws[f"K{row}"].value != TEST_RATE:
            fail(f"{ccy} rate cell does not hold the feed's rate")
        if f"$K${row}" not in ws[f"L{row}"].value or f"$K${row}" not in ws[f"M{row}"].value:
            fail(f"{ccy} → HKD does not use the rate cell")
        cost, gain = sums(ccy)
        cost, gain = cost * TEST_RATE, gain * TEST_RATE
        if not (_close(sheet.value(f"L{row}"), cost) and _close(sheet.value(f"M{row}"), gain)):
            fail(f"{ccy} → HKD is not the {ccy} rows times the rate")
        if row <= max(r for r, c in details if c == ccy):
            fail(f"{ccy} → HKD sits above its own rows")
        totals.append((ccy, row, cost, gain))

    grand = [row for label, row in labels.items() if label.endswith("Grand Total")]
    totals.sort(key=lambda t: t[1])
    if len(totals) > 1:
        want = " and ".join([", ".join(c for c, *_ in totals[:-1]), totals[-1][0]]) + " Grand Total"
        if want not in labels:
            fail(f"no '{want}'")
        else:
            row = labels[want]
            cost, gain = sum(t[2] for t in totals), sum(t[3] for t in totals)
            if not (_close(sheet.value(f"L{row}"), cost) and _close(sheet.value(f"M{row}"), gain)):
                fail("Grand Total is not the HKD Sub Total plus the converted rows")
    elif grand:
        fail("Grand Total written with a single total")

    # Feed down: no figure may look complete without the foreign currency.
    ws, _ = _generate(db, trade_date, summary, None)
    sheet = Sheet(ws)
    _, _, labels = _layout(ws)
    for ccy in foreign:
        row = labels.get(f"{ccy} → HKD")
        if row is None or ws[f"K{row}"].value != "enter rate":
            fail(f"feed down: {ccy} rate cell is not marked 'enter rate'")
        elif sheet.safe(f"L{row}") != "#VALUE!":
            fail(f"feed down: {ccy} → HKD shows a figure")
    for label, row in labels.items():
        if label.endswith("Grand Total") and sheet.safe(f"L{row}") != "#VALUE!":
            fail("feed down: Grand Total shows a figure")
    return ok


def _check_unchanged(db, trade_date, summary):
    ws, calls = _generate(db, trade_date, summary, TEST_RATE)
    _, _, labels = _layout(ws)
    if calls or labels:
        print(f"  {trade_date}: FAIL  totals or a rate fetch on a day with no other-currency gain/loss")
        return False
    return True


def _check_feed():
    ok = True
    original = module.urlopen

    def reply(body):
        return lambda url, timeout: io.BytesIO(json.dumps(body).encode())

    def down(url, timeout):
        raise OSError("network down")

    cases = [
        (down, None, "network error"),
        (reply({"result": "error", "error-type": "unsupported-code"}), None, "error reply"),
        (reply({"result": "success", "rates": {"USD": 1}}), None, "no HKD rate"),
        (reply({"result": "success", "rates": {"HKD": 6.123456}, "time_last_update_utc": "t"}),
         (6.1235, "t"), "good reply"),
    ]
    try:
        for fake, want, name in cases:
            module.urlopen = fake
            got = module.fetch_hkd_rate("SGD")
            if got != want:
                print(f"  feed, {name}: FAIL  got {got!r}")
                ok = False
    finally:
        module.urlopen = original
    return ok


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    db_path = sys.argv[1] if len(sys.argv) > 1 else DB_PATH
    ctx, db = _open(db_path)
    real_feed = module.fetch_hkd_rate
    try:
        ok = _check_feed()
        checked, unchanged = [], 0
        for d in _dates(db):
            summary = TransactionSummary(db, d)
            if summary.is_empty():
                continue
            foreign = _foreign_gain(summary)
            if foreign:
                ok &= _check(db, d, summary, foreign)
                checked.append(d)
            else:
                ok &= _check_unchanged(db, d, summary)
                unchanged += 1
        if not checked:
            print("  FAIL: no date with a gain/loss section in another currency")
            ok = False
        print(f"  {len(checked)} date(s) with another currency checked: {', '.join(checked)}")
        print(f"  {unchanged} other date(s): no totals, no rate fetched")
    finally:
        module.fetch_hkd_rate = real_feed
        ctx.pop()

    print("RESULT:", "ALL PASS" if ok else "FAIL")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
